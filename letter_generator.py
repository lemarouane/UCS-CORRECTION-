"""
Generation de la demande d'ajustement / lettre (cahier, section 14.3).

Le cahier precise que le texte final doit etre fourni/valide par le metier.
En attendant ce validation, ce module genere un document Word generique et
STATIQUE qui fonctionne pour tous les cas : il reprend les informations de
l'ajustement (projet, date, version du modele) et la liste des
modifications de longueur, sans texte commercial/juridique invente.
A remplacer/adapter une fois le template metier fourni.
"""

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import MODEL_VERSION


def generate_adjustment_letter(results_df, project_name=""):
    """
    Genere une lettre/demande d'ajustement au format .docx a partir du
    tableau de resultats.

    :param results_df: DataFrame produit par processor.process_ucs_file
    :param project_name: nom du projet/faisceau a afficher en en-tete (optionnel)
    :return: bytes du document .docx
    """
    modified_df = results_df[
        (results_df["Delta arrondi (mm)"].notna()) & (results_df["Delta arrondi (mm)"] != 0)
    ].copy()

    document = Document()
    _set_base_style(document)

    title = document.add_heading("Demande d'ajustement de longueurs", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info_paragraph = document.add_paragraph()
    info_paragraph.add_run("Projet / Faisceau : ").bold = True
    info_paragraph.add_run(project_name or "-")

    date_paragraph = document.add_paragraph()
    date_paragraph.add_run("Date : ").bold = True
    date_paragraph.add_run(datetime.now().strftime("%d/%m/%Y %H:%M"))

    version_paragraph = document.add_paragraph()
    version_paragraph.add_run("Version du modele : ").bold = True
    version_paragraph.add_run(MODEL_VERSION)

    count_paragraph = document.add_paragraph()
    count_paragraph.add_run("Nombre d'elements modifies : ").bold = True
    count_paragraph.add_run(str(len(modified_df)))

    document.add_paragraph(
        "Le tableau ci-dessous liste les Wire/Tube dont la longueur doit "
        "etre ajustee selon l'estimation automatique, arrondie au pas "
        "metier de 5 mm. Cette estimation est une aide au calcul et ne "
        "constitue pas une garantie universelle."
    )

    _add_modifications_table(document, modified_df)

    document.add_paragraph()
    signature_paragraph = document.add_paragraph()
    signature_paragraph.add_run("Validation metier : ").bold = True
    signature_paragraph.add_run("_______________________")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _set_base_style(document):
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _add_modifications_table(document, modified_df):
    headers = ["Feuille", "Wire/Tube Name", "Old Length (mm)", "Delta arrondi (mm)", "New Length (mm)"]

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"

    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        header_cells[idx].text = header
        for paragraph in header_cells[idx].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for _, record in modified_df.iterrows():
        row_cells = table.add_row().cells
        row_cells[0].text = str(record["Feuille"])
        row_cells[1].text = str(record["Wire/Tube Name"])
        row_cells[2].text = str(record["Old Length (mm)"])
        row_cells[3].text = str(record["Delta arrondi (mm)"])
        row_cells[4].text = str(record["New Length (mm)"])

    if len(modified_df) == 0:
        no_row = table.add_row().cells
        no_row[0].merge(no_row[-1])
        no_row[0].text = "Aucune modification a signaler."