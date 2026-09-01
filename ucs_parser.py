"""
Lecture et normalisation d'un fichier UCS (.xlsx).

Un UCS est un classeur Excel avec (au minimum) une feuille "Wires"
contenant une colonne de longueur ("Length") et une colonne de nom
("Wire Name" / "Wire name" / "Name"). L'en-tete n'est pas forcement sur la
ligne 1 (certains exports ont une ligne de version en ligne 1, ex :
"EngApp v3.1.0").

Les Tubes sont volontairement hors perimetre (cf. config.SHEETS_TO_PROCESS) :
seuls les Wires sont pris en compte par le calcul automatique.

Ce module NE MODIFIE JAMAIS le fichier source : il ne fait que lire.
"""

from dataclasses import dataclass

import openpyxl

from config import (
    SHEETS_TO_PROCESS,
    LENGTH_HEADER_KEYWORDS,
    NAME_HEADER_INCLUDE_KEYWORDS,
    NAME_HEADER_EXCLUDE_KEYWORDS,
)


@dataclass
class SheetLayout:
    sheet_name: str
    header_row_index: int  # 1-based, ligne des en-tetes
    name_col_index: int  # 1-based
    length_col_index: int  # 1-based
    name_header: str
    length_header: str


@dataclass
class ParsedRow:
    sheet_name: str
    excel_row_index: int  # 1-based, position reelle dans le classeur (pour re-ecriture)
    wire_name: object
    old_length_raw: object


class UCSStructureError(Exception):
    """Levee lorsque la structure UCS est inconnue ou non conforme."""


def _normalize_header(value):
    return str(value).strip().lower() if value is not None else ""


def _find_header_row_and_columns(worksheet, max_scan_rows=10):
    """
    Cherche, dans les `max_scan_rows` premieres lignes, la ligne d'en-tete
    contenant une colonne "Length" et une colonne "Name" exploitable.
    """
    for row_idx in range(1, max_scan_rows + 1):
        row_values = [cell.value for cell in worksheet[row_idx]]
        normalized = [_normalize_header(v) for v in row_values]

        length_col = None
        for col_idx, header in enumerate(normalized, start=1):
            if header in LENGTH_HEADER_KEYWORDS:
                length_col = col_idx
                break
        if length_col is None:
            continue

        name_col = None
        for col_idx, header in enumerate(normalized, start=1):
            if any(inc in header for inc in NAME_HEADER_INCLUDE_KEYWORDS) and not any(
                exc in header for exc in NAME_HEADER_EXCLUDE_KEYWORDS
            ):
                name_col = col_idx
                break

        if name_col is not None:
            return row_idx, name_col, length_col, row_values[name_col - 1], row_values[length_col - 1]

    return None, None, None, None, None


def detect_sheet_layout(workbook, sheet_name):
    """
    Detecte automatiquement la ligne d'en-tete et les colonnes Name/Length
    d'une feuille. Retourne None si la feuille n'existe pas ou n'est pas
    exploitable.
    """
    if sheet_name not in workbook.sheetnames:
        return None

    worksheet = workbook[sheet_name]
    header_row, name_col, length_col, name_header, length_header = _find_header_row_and_columns(worksheet)

    if header_row is None:
        return None

    return SheetLayout(
        sheet_name=sheet_name,
        header_row_index=header_row,
        name_col_index=name_col,
        length_col_index=length_col,
        name_header=str(name_header) if name_header is not None else "",
        length_header=str(length_header) if length_header is not None else "",
    )


def parse_ucs_file(file_path_or_buffer, sheets_to_process=None):
    """
    Parse un fichier UCS et retourne :
      - layouts: dict {sheet_name: SheetLayout} pour les feuilles detectees
      - rows: list[ParsedRow] pour toutes les feuilles detectees
      - skipped_sheets: list[str] des feuilles demandees mais non detectees

    Ne leve UCSStructureError que si AUCUNE feuille attendue n'a pu etre
    detectee (cahier, section 12 : ne pas modifier le fichier source si la
    structure est inconnue).
    """
    sheets_to_process = sheets_to_process or SHEETS_TO_PROCESS

    workbook = openpyxl.load_workbook(file_path_or_buffer, data_only=True)

    layouts = {}
    rows = []
    skipped_sheets = []

    for sheet_name in sheets_to_process:
        layout = detect_sheet_layout(workbook, sheet_name)
        if layout is None:
            skipped_sheets.append(sheet_name)
            continue

        layouts[sheet_name] = layout
        worksheet = workbook[sheet_name]

        for row_idx in range(layout.header_row_index + 1, worksheet.max_row + 1):
            name_cell = worksheet.cell(row=row_idx, column=layout.name_col_index)
            length_cell = worksheet.cell(row=row_idx, column=layout.length_col_index)

            if name_cell.value is None and length_cell.value is None:
                continue  # ligne vide

            rows.append(
                ParsedRow(
                    sheet_name=sheet_name,
                    excel_row_index=row_idx,
                    wire_name=name_cell.value,
                    old_length_raw=length_cell.value,
                )
            )

    if not layouts:
        raise UCSStructureError(
            "Structure UCS inconnue ou non conforme : aucune feuille parmi "
            f"{sheets_to_process} n'a pu etre detectee avec des colonnes "
            "Name/Length exploitables."
        )

    return layouts, rows, skipped_sheets